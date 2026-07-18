"""
Calmora - Export Utilities
--------------------------
Generic helpers that turn a list of dicts into a downloadable CSV, XLSX,
or PDF file. Each module's export route builds a list[dict] of the rows
it wants to export, then calls one of these three functions.
"""

import io
import csv
from datetime import datetime

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch


def export_csv(rows: list[dict], filename: str):
    """rows: list of dicts with identical keys -> CSV file download."""
    if not rows:
        rows = [{"info": "No data available"}]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
    writer.writeheader()
    writer.writerows(rows)

    mem = io.BytesIO(buffer.getvalue().encode("utf-8"))
    mem.seek(0)
    return send_file(
        mem,
        mimetype="text/csv",
        as_attachment=True,
        download_name=f"{filename}_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
    )


def export_excel(rows: list[dict], filename: str, sheet_title: str = "Data"):
    if not rows:
        rows = [{"info": "No data available"}]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers = list(rows[0].keys())
    ws.append([h.replace("_", " ").title() for h in headers])
    header_fill = PatternFill(start_color="6C63FF", end_color="6C63FF", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(
        mem,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"{filename}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
    )


def export_pdf(rows: list[dict], filename: str, title: str = "Calmora Report"):
    if not rows:
        rows = [{"info": "No data available"}]
    mem = io.BytesIO()
    doc = SimpleDocTemplate(mem, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    elements.append(
        Paragraph(f"Generated {datetime.now().strftime('%B %d, %Y %H:%M')}", styles["Normal"])
    )
    elements.append(Spacer(1, 16))

    headers = list(rows[0].keys())
    table_data = [[h.replace("_", " ").title() for h in headers]]
    for row in rows:
        table_data.append([str(row.get(h, "")) for h in headers])

    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6C63FF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    mem.seek(0)
    return send_file(
        mem,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"{filename}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
    )
